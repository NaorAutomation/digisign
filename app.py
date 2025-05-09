from flask import Flask, request, render_template, redirect, url_for, flash, session, jsonify
import os
from datetime import datetime
import uuid
import db
import storage
from flask_cors import CORS
from pdfminer.high_level import extract_text
import tempfile
import docx2txt
import re
from docx import Document
import html

app = Flask(__name__)
CORS(app)  # הוספת תמיכה ב-CORS
app.secret_key = os.environ.get('FLASK_SECRET_KEY', 'digital-signature-app-secret')

# פונקציה לחילוץ טקסט מקובץ DOCX עם שמירה על מבנה
def extract_text_from_docx_with_structure(docx_file):
    try:
        doc = Document(docx_file)
        result = []
        
        # עיבוד פסקאות
        for para in doc.paragraphs:
            if not para.text.strip():
                # הוספת שורה ריקה
                result.append("<p>&nbsp;</p>")
                continue
                
            # בדיקה האם זה כותרת או פסקה רגילה
            if para.style.name.startswith('Heading'):
                # בדיקה לרמת הכותרת (Heading 1, Heading 2, וכו')
                heading_level = int(para.style.name.split()[-1]) if para.style.name.split()[-1].isdigit() else 1
                result.append(f"<h{heading_level}>{html.escape(para.text)}</h{heading_level}>")
            else:
                # טיפול בסעיפים ממוספרים
                text = para.text
                
                # טיפול במספרים וסעיפים - זיהוי תבניות שכיחות
                # זיהוי סעיפים מספריים (1. 2. וכו')
                if re.match(r'^\d+[\.\)](\s|\u00A0)', text):
                    result.append(f"<p class='numbered-paragraph section-numeric'>{html.escape(text)}</p>")
                # זיהוי סעיפים עם אותיות עבריות (א. ב. וכו')
                elif re.match(r'^[\u0590-\u05FF][\.\)](\s|\u00A0)', text):
                    result.append(f"<p class='numbered-paragraph section-letter'>{html.escape(text)}</p>")
                # זיהוי סעיפי משנה בעזרת מקף או תבליט
                elif text.strip().startswith('-') or text.strip().startswith('•'):
                    result.append(f"<p class='numbered-paragraph subsection'>{html.escape(text)}</p>")
                else:
                    # פסקה רגילה
                    result.append(f"<p>{html.escape(text)}</p>")
        
        # טיפול בטבלאות
        for table in doc.tables:
            table_html = "<table class='extracted-table' style='width:100%; border-collapse:collapse; margin:10px 0;'>"
            
            # מעבר על כל שורה בטבלה
            for row in table.rows:
                table_html += "<tr>"
                # מעבר על כל תא בשורה
                for cell in row.cells:
                    table_html += "<td style='border:1px solid #ddd; padding:8px;'>"
                    # הוספת תוכן התא (כל הפסקאות בתא)
                    for paragraph in cell.paragraphs:
                        table_html += f"<p>{html.escape(paragraph.text)}</p>"
                    table_html += "</td>"
                table_html += "</tr>"
            
            table_html += "</table>"
            result.append(table_html)
        
        # חיבור כל החלקים לקוד HTML אחד
        html_content = "\n".join(result)
        
        return html_content
    except Exception as e:
        print(f"שגיאה בחילוץ מבנה מסמך DOCX: {str(e)}")
        # במקרה של שגיאה, חזרה לשיטה הרגילה
        return docx2txt.process(docx_file)

# פונקציה לחילוץ טקסט מתוך קובץ
def extract_text_from_file(file):
    # שמירת הקובץ למיקום זמני
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(file.filename)[1])
    file.save(temp_file.name)
    temp_file.close()
    
    file_ext = os.path.splitext(file.filename)[1].lower()
    print(f"מנסה לחלץ טקסט מקובץ: {file.filename}, סוג: {file_ext}")
    
    try:
        text = ""
        
        # טיפול בקבצי PDF
        if file_ext == '.pdf':
            print("מחלץ טקסט מקובץ PDF...")
            text = extract_text(temp_file.name)
            
            # המרת טקסט PDF למבנה HTML בסיסי - שמירה על שורות ופסקאות
            if text:
                text_with_structure = []
                paragraphs = text.split('\n\n')
                
                for paragraph in paragraphs:
                    if not paragraph.strip():
                        continue
                        
                    # בדיקה אם מדובר בכותרת
                    if len(paragraph.strip()) > 0 and len(paragraph.strip()) < 100 and paragraph.strip().isupper():
                        text_with_structure.append(f"<h2>{html.escape(paragraph.strip())}</h2>")
                    # זיהוי סעיפים מספריים בעברית
                    elif re.match(r'^\d+[\.\)](\s|\u00A0)', paragraph.strip()):
                        text_with_structure.append(f"<p class='numbered-paragraph section-numeric'>{html.escape(paragraph.strip())}</p>")
                    # זיהוי סעיפים עם אותיות עבריות
                    elif re.match(r'^[\u0590-\u05FF][\.\)](\s|\u00A0)', paragraph.strip()):
                        text_with_structure.append(f"<p class='numbered-paragraph section-letter'>{html.escape(paragraph.strip())}</p>")
                    # זיהוי סעיפי משנה
                    elif paragraph.strip().startswith('-') or paragraph.strip().startswith('•'):
                        text_with_structure.append(f"<p class='numbered-paragraph subsection'>{html.escape(paragraph.strip())}</p>")
                    else:
                        # כל השאר פסקאות רגילות
                        lines = paragraph.split('\n')
                        paragraph_html = "<p>" + html.escape(lines[0])
                        
                        for line in lines[1:]:
                            if line.strip():
                                paragraph_html += "<br>" + html.escape(line)
                                
                        paragraph_html += "</p>"
                        text_with_structure.append(paragraph_html)
                
                text = "\n".join(text_with_structure)
        
        # טיפול בקבצי DOCX
        elif file_ext == '.docx':
            print("מחלץ טקסט מקובץ DOCX עם שמירה על מבנה...")
            try:
                # שימוש בפונקציה המשופרת לחילוץ עם מבנה
                text = extract_text_from_docx_with_structure(temp_file.name)
                print(f"הצלחה בחילוץ טקסט משופר, אורך התוכן: {len(text)}")
            except Exception as docx_error:
                print(f"שגיאה ספציפית בעיבוד DOCX: {str(docx_error)}")
                # ניסיון לפעול בשיטה הרגילה
                try:
                    text = docx2txt.process(temp_file.name)
                    print(f"אורך הטקסט שחולץ בשיטה רגילה: {len(text)}")
                except Exception as docx2_error:
                    print(f"שגיאה גם עם docx2txt: {str(docx2_error)}")
                    text = "שגיאה בחילוץ הטקסט מקובץ DOCX. נא להזין את הטקסט ידנית."
        
        # טיפול בקבצי DOC
        elif file_ext == '.doc':
            print("מחלץ טקסט מקובץ DOC...")
            # כרגע נחזיר הודעה - יש להוסיף תמיכה מלאה בהמשך
            text = "קובץ DOC הועלה בהצלחה. נא להעתיק את הטקסט לכאן ידנית."
        
        # טיפול בקבצי טקסט פשוטים
        elif file_ext in ['.txt', '.text']:
            print("מחלץ טקסט מקובץ TXT...")
            with open(temp_file.name, 'r', encoding='utf-8', errors='ignore') as f:
                raw_text = f.read()
                
                # המרת הטקסט למבנה HTML בסיסי
                paragraphs = raw_text.split('\n\n')
                text_with_structure = []
                
                for paragraph in paragraphs:
                    if not paragraph.strip():
                        continue
                        
                    # בדיקה אם מדובר בסעיף ממוספר
                    if re.match(r'^\d+[\.\)]\s', paragraph.strip()) or re.match(r'^[\u0590-\u05FF]+[\.\)]\s', paragraph.strip()):
                        text_with_structure.append(f"<p class='numbered-paragraph'>{html.escape(paragraph.strip())}</p>")
                    else:
                        # החלפת שורות בודדות בתגי <br>
                        paragraph_html = "<p>" + html.escape(paragraph.replace('\n', '<br>')) + "</p>"
                        text_with_structure.append(paragraph_html)
                
                text = "\n".join(text_with_structure)
        
        # טיפול בקבצי RTF
        elif file_ext == '.rtf':
            print("מחלץ טקסט מקובץ RTF...")
            # כרגע נחזיר הודעה - יש להוסיף תמיכה מלאה בהמשך
            text = "קובץ RTF הועלה בהצלחה. נא להעתיק את הטקסט לכאן ידנית."
        
        # פורמטים לא נתמכים
        else:
            print(f"פורמט לא נתמך: {file_ext}")
            text = f"הפורמט {file_ext} אינו נתמך לחילוץ טקסט אוטומטי. נא להזין את הטקסט ידנית."
        
        # בדיקה אם הטקסט ריק
        if not text or text.strip() == "" or len(text.strip()) < 10:
            print("הטקסט שחולץ קצר מדי או ריק")
            text = "<p>לא נמצא טקסט בקובץ זה או שהקובץ מוגן. נא לנסות להזין טקסט באופן ידני.</p>"
            
    except Exception as e:
        print(f"שגיאה כללית בחילוץ טקסט: {str(e)}")
        text = f"<p>שגיאה בחילוץ טקסט: {str(e)}</p>"
    finally:
        # ניקוי קבצים זמניים
        if os.path.exists(temp_file.name):
            print(f"מוחק קובץ זמני: {temp_file.name}")
            os.unlink(temp_file.name)
        else:
            print(f"קובץ זמני לא נמצא: {temp_file.name}")
    
    return text

# הוספת פונקציית עזר לתאריך שתהיה זמינה בתבניות
@app.context_processor
def utility_processor():
    def now():
        return datetime.now()
    return {"now": now}

# עמוד הבית
@app.route('/')
def index():
    return render_template('index.html')

# אימות עסק לפי ח.פ
@app.route('/business/auth', methods=['POST'])
def business_auth():
    business_id_number = request.form.get('business_id_number')
    business = db.get_business_by_id_number(business_id_number)
    
    if business:
        session['business_id'] = business['business_id']
        return redirect(url_for('business_dashboard'))
    else:
        session['temp_business_id_number'] = business_id_number
        return redirect(url_for('business_register'))

# רישום עסק חדש
@app.route('/business/register', methods=['GET', 'POST'])
def business_register():
    if request.method == 'GET':
        business_id_number = session.get('temp_business_id_number', '')
        return render_template('business_register.html', business_id_number=business_id_number)
    
    if request.method == 'POST':
        business_data = {
            'business_id_number': request.form.get('business_id_number'),
            'business_name': request.form.get('business_name'),
            'contact_name': request.form.get('contact_name'),
            'email': request.form.get('email'),
            'phone': request.form.get('phone'),
            'address': request.form.get('address')
        }
        
        if 'logo' in request.files:
            logo_file = request.files['logo']
            if logo_file.filename != '':
                logo_path = storage.upload_logo(logo_file)
                business_data['logo_path'] = logo_path
        
        business_id = db.create_business(business_data)
        session['business_id'] = business_id
        
        flash('העסק נרשם בהצלחה!', 'success')
        return redirect(url_for('business_dashboard'))

# דף ניהול העסק
@app.route('/business/dashboard')
def business_dashboard():
    if 'business_id' not in session:
        return redirect(url_for('index'))
    
    business_id = session['business_id']
    business = db.get_business(business_id)
    contracts = db.get_business_contracts(business_id)
    
    # עיבוד נתוני תאריך כדי למנוע שגיאת strftime
    for contract in contracts:
        if isinstance(contract['created_at'], str):
            # אם התאריך כבר מחרוזת, נשאיר אותו כמות שהוא
            created_date = contract['created_at']
        else:
            # אם התאריך הוא אובייקט datetime, נמיר אותו למחרוזת מפורמטת
            try:
                created_date = contract['created_at'].strftime('%d/%m/%Y %H:%M')
            except (AttributeError, TypeError):
                created_date = "לא צוין"
        
        contract['created_at'] = created_date
    
    return render_template('business_dashboard.html', business=business, contracts=contracts)

# יצירת חוזה חדש
@app.route('/contract/create', methods=['GET', 'POST'])
def create_contract():
    if 'business_id' not in session:
        return redirect(url_for('index'))
    
    business_id = session['business_id']
    
    if request.method == 'GET':
        return render_template('create_contract.html')
    
    if request.method == 'POST':
        # נתונים בסיסיים של החוזה
        contract_data = {
            'business_id': business_id,
            'status': 'pending',
            'notes': request.form.get('notes'),
            'customer_name': request.form.get('customer_name'),
            'customer_email': request.form.get('customer_email'),
            'customer_phone': request.form.get('customer_phone')
        }
        
        print(f"יוצר חוזה חדש: טקסט={bool(request.form.get('text_content'))}, קובץ={bool('document' in request.files and request.files['document'].filename != '')}")
        
        # בדיקה אם יש טקסט או קובץ
        if 'document' in request.files and request.files['document'].filename != '':
            # יש קובץ - מצב 'file'
            document_file = request.files['document']
            contract_data['contract_type'] = 'file'
            file_path = storage.upload_document(document_file)
            contract_data['pdf_path'] = file_path
            print(f"קובץ נשמר בנתיב: {file_path}")
        else:
            # אין קובץ, צריך להיות טקסט - מצב 'text'
            contract_data['contract_type'] = 'text'
            contract_data['text_content'] = request.form.get('text_content')
            print(f"נשמר תוכן טקסט באורך: {len(contract_data['text_content'] or '')}")
        
        # יצירת URL ייחודי
        unique_url = str(uuid.uuid4())[:12]
        contract_data['unique_url'] = unique_url
        
        contract_id = db.create_contract(contract_data)
        
        # יצירת קישור WhatsApp אם יש טלפון
        if contract_data['customer_phone']:
            customer_phone = contract_data['customer_phone'].replace('-', '').replace(' ', '')
            contract_url = request.host_url + 'contract/' + unique_url
            whatsapp_url = f"https://wa.me/{customer_phone}?text=אנא חתום על החוזה בקישור הבא: {contract_url}"
            return render_template('contract_sent.html', contract_id=contract_id, whatsapp_url=whatsapp_url)
        
        flash('החוזה נוצר בהצלחה!', 'success')
        return redirect(url_for('business_dashboard'))

# צפייה בחוזה ללקוח
@app.route('/contract/<unique_url>')
def view_contract(unique_url):
    contract = db.get_contract_by_url(unique_url)
    
    if not contract:
        flash('החוזה לא נמצא.', 'error')
        return redirect(url_for('index'))
    
    business = db.get_business(contract['business_id'])
    
    # עיבוד מקדים של נתיבי התמונות לקבלת URLs מלאים
    if contract.get('customer_id_path'):
        # בדיקה אם הנתיב כבר מכיל URL מלא
        if not contract['customer_id_path'].startswith('http'):
            try:
                # קבלת URL ציבורי לתעודת זהות
                contract['customer_id_path'] = storage.get_public_url(contract['customer_id_path'])
            except Exception as e:
                print(f"שגיאה בהמרת נתיב תעודת זהות ל-URL: {str(e)}")
    
    if contract.get('signature_path'):
        # בדיקה אם הנתיב כבר מכיל URL מלא
        if not contract['signature_path'].startswith('http'):
            try:
                # קבלת URL ציבורי לחתימה
                contract['signature_path'] = storage.get_public_url(contract['signature_path'])
            except Exception as e:
                print(f"שגיאה בהמרת נתיב חתימה ל-URL: {str(e)}")
    
    return render_template('customer_view.html', contract=contract, business=business)

# חתימה על חוזה
@app.route('/contract/<unique_url>/sign', methods=['POST'])
def sign_contract(unique_url):
    contract = db.get_contract_by_url(unique_url)
    
    if not contract:
        return jsonify({'error': 'החוזה לא נמצא'}), 404
    
    # עדכון פרטי לקוח
    customer_data = {}
    if request.form.get('customer_name'):
        customer_data['customer_name'] = request.form.get('customer_name')
    if request.form.get('customer_email'):
        customer_data['customer_email'] = request.form.get('customer_email')
    if request.form.get('customer_phone'):
        customer_data['customer_phone'] = request.form.get('customer_phone')
    
    print("התקבלו קבצים:", request.files.keys())
    
    # העלאת ת.ז
    if 'customer_id' in request.files and request.files['customer_id'].filename != '':
        try:
            id_file = request.files['customer_id']
            print(f"קובץ ת.ז התקבל: {id_file.filename}")
            id_path = storage.upload_id(id_file)
            customer_data['customer_id_path'] = id_path
            print(f"נתיב ת.ז נשמר: {id_path}")
        except Exception as e:
            print(f"שגיאה בהעלאת קובץ ת.ז: {str(e)}")
    else:
        print("לא התקבל קובץ ת.ז")
    
    # טיפול בחתימה
    if 'signature' in request.files and request.files['signature'].filename != '':
        try:
            signature_file = request.files['signature']
            print(f"קובץ חתימה התקבל: {signature_file.filename}")
            signature_path = storage.upload_signature(signature_file)
            customer_data['signature_path'] = signature_path
            # תיעוד זמן החתימה
            customer_data['signature_timestamp'] = datetime.now().isoformat()
            # עדכון סטטוס
            customer_data['status'] = 'signed'
            print(f"נתיב חתימה נשמר: {signature_path}")
        except Exception as e:
            print(f"שגיאה בהעלאת קובץ חתימה: {str(e)}")
    else:
        print("לא התקבל קובץ חתימה")
    
    print("מעדכן חוזה עם נתונים:", customer_data)
    db.update_contract(contract['contract_id'], customer_data)
    
    return jsonify({'success': True})

# נקודת קצה לחילוץ טקסט מקובץ
@app.route('/api/extract-text', methods=['POST'])
def api_extract_text():
    if 'file' not in request.files:
        return jsonify({'error': 'לא התקבל קובץ'}), 400
    
    uploaded_file = request.files['file']
    if uploaded_file.filename == '':
        return jsonify({'error': 'לא נבחר קובץ'}), 400
    
    # בדיקה שהקובץ בפורמט מאושר
    file_ext = os.path.splitext(uploaded_file.filename)[1].lower()
    allowed_extensions = ['.pdf', '.docx', '.doc', '.txt', '.text', '.rtf']
    
    if file_ext not in allowed_extensions:
        return jsonify({'error': f'סוג הקובץ אינו נתמך. הפורמטים הנתמכים: {", ".join(allowed_extensions)}'}), 400
    
    try:
        extracted_text = extract_text_from_file(uploaded_file)
        print(f"הטקסט חולץ בהצלחה. אורך הטקסט: {len(extracted_text)}")
        
        # מידע נוסף על הקובץ המקורי
        file_info = {
            'originalName': uploaded_file.filename,
            'fileType': file_ext.replace('.', ''),
            'timestamp': datetime.now().isoformat()
        }
        
        return jsonify({
            'text': extracted_text,
            'fileInfo': file_info,
            'success': True
        })
    except Exception as e:
        print(f"שגיאה לא צפויה בעת חילוץ הטקסט: {str(e)}")
        return jsonify({'error': f'שגיאה בעת חילוץ הטקסט: {str(e)}', 'text': 'אירעה שגיאה בעת חילוץ הטקסט. אנא נסה שוב או הזן את הטקסט ידנית.'}), 500

if __name__ == '__main__':
    app.run(debug=True)